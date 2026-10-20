# -*- coding: utf-8 -*-
"""
Генератор отчетов по Лабораторным работам №6 и №8 для СПБ ГУАП (Кафедра №43).
Создает документы Microsoft Word (.docx) и компилирует их в PDF (.pdf) через Microsoft Word.
"""

import os
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from docx2pdf import convert

def set_font(run, name="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    r = run._r
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rPr.insert(0, rFonts)

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False,
          size=12, space_before=0, space_after=4, first_indent=1.25):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if first_indent > 0:
        pf.first_line_indent = Cm(first_indent)
    else:
        pf.first_line_indent = Pt(0)
    if text:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, italic=italic)
    return p

def add_heading(doc, text, level=1):
    sizes = {1: 14, 2: 12, 3: 12}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(4)
    pf.first_line_indent = Pt(0)
    run = p.add_run(text)
    set_font(run, size=sizes.get(level, 12), bold=True)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Cm(1.25)
    run = p.add_run(text)
    set_font(run, size=12)
    return p

def add_code(doc, code_str):
    for line in code_str.strip().split("\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.left_indent = Cm(0.8)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        run = p.add_run(line if line else " ")
        set_font(run, name="Courier New", size=9.5)
    add_p(doc, "", first_indent=0, space_before=0, space_after=4)

def set_cell_border(cell, **kwargs):
    """
    kwargs: bottom, top, left, right.
    val: "single", "nil", "none", etc.
    sz: "4", "8", etc.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                          f'<w:top w:val="{kwargs.get("top", "none")}" w:sz="{kwargs.get("top_sz", "4")}" w:space="0" w:color="auto"/>'
                          f'<w:left w:val="{kwargs.get("left", "none")}" w:sz="{kwargs.get("left_sz", "4")}" w:space="0" w:color="auto"/>'
                          f'<w:bottom w:val="{kwargs.get("bottom", "none")}" w:sz="{kwargs.get("bottom_sz", "4")}" w:space="0" w:color="auto"/>'
                          f'<w:right w:val="{kwargs.get("right", "none")}" w:sz="{kwargs.get("right_sz", "4")}" w:space="0" w:color="auto"/>'
                          f'</w:tcBorders>')
    tcPr.append(tcBorders)

def build_title_page(doc, lab_num, lab_title):
    # Заголовок ВУЗа
    p1 = add_p(doc, "ГУАП", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, first_indent=0, space_after=2)
    p2 = add_p(doc, "КАФЕДРА № 43", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, first_indent=0, space_after=14)

    # ОТЧЕТ ЗАЩИЩЕН С ОЦЕНКОЙ
    p3 = add_p(doc, "ОТЧЕТ", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=12, first_indent=0, space_after=1)
    p4 = add_p(doc, "ЗАЩИЩЕН С ОЦЕНКОЙ", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=12, first_indent=0, space_after=6)
    p5 = add_p(doc, "ПРЕПОДАВАТЕЛЬ", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=12, first_indent=0, space_after=2)

    # Таблица преподавателя (3 колонки: должность, подпись, ФИО)
    t_prep = doc.add_table(rows=2, cols=3)
    t_prep.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(5.5), Cm(4.0), Cm(5.5)]
    for row in t_prep.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width

    # Верхний ряд
    c0 = t_prep.cell(0, 0)
    p_c0 = c0.paragraphs[0]
    p_c0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c0 = p_c0.add_run("Доцент, канд. тех. наук")
    set_font(r_c0, size=11)
    set_cell_border(c0, bottom="single", bottom_sz="6")

    c1 = t_prep.cell(0, 1)
    p_c1 = c1.paragraphs[0]
    p_c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_border(c1, bottom="single", bottom_sz="6")

    c2 = t_prep.cell(0, 2)
    p_c2 = c2.paragraphs[0]
    p_c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c2 = p_c2.add_run("Степанов П. А.")
    set_font(r_c2, size=11)
    set_cell_border(c2, bottom="single", bottom_sz="6")

    # Нижний ряд (подписи под чертой)
    sub_labels = ["должность, уч. степень, звание", "подпись, дата", "инициалы, фамилия"]
    for i, lbl in enumerate(sub_labels):
        cell = t_prep.cell(1, i)
        p_sub = cell.paragraphs[0]
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(1)
        r_sub = p_sub.add_run(lbl)
        set_font(r_sub, size=9)

    add_p(doc, "", first_indent=0, space_after=14)

    # Центральный блок: Отчет
    add_p(doc, f"ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ №{lab_num}", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, first_indent=0, space_after=4)
    add_p(doc, lab_title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, first_indent=0, space_after=6)
    add_p(doc, "по курсу: Технология разработки серверных информационных систем", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, first_indent=0, space_after=24)

    # РАБОТУ ВЫПОЛНИЛ
    add_p(doc, "РАБОТУ ВЫПОЛНИЛ", align=WD_ALIGN_PARAGRAPH.LEFT, size=12, first_indent=0, space_after=4)

    # Таблица студента
    t_stud = doc.add_table(rows=2, cols=3)
    t_stud.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in t_stud.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width

    sc0 = t_stud.cell(0, 0)
    p_sc0 = sc0.paragraphs[0]
    p_sc0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sc0 = p_sc0.add_run("СТУДЕНТ гр. № 4434")
    set_font(r_sc0, size=11)
    set_cell_border(sc0, bottom="single", bottom_sz="6")

    sc1 = t_stud.cell(0, 1)
    p_sc1 = sc1.paragraphs[0]
    p_sc1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_border(sc1, bottom="single", bottom_sz="6")

    sc2 = t_stud.cell(0, 2)
    p_sc2 = sc2.paragraphs[0]
    p_sc2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sc2 = p_sc2.add_run("Энгельгардт К. А.")
    set_font(r_sc2, size=11)
    set_cell_border(sc2, bottom="single", bottom_sz="6")

    stud_sub_labels = ["", "подпись, дата", "инициалы, фамилия"]
    for i, lbl in enumerate(stud_sub_labels):
        cell = t_stud.cell(1, i)
        p_sub = cell.paragraphs[0]
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(1)
        r_sub = p_sub.add_run(lbl)
        set_font(r_sub, size=9)

    add_p(doc, "", first_indent=0, space_after=32)

    # Город и год
    add_p(doc, "Санкт-Петербург 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, first_indent=0, space_after=0)
    doc.add_page_break()

# ==============================================================================
# СБОРКА ОТЧЕТА №6
# ==============================================================================
def generate_report_6():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(1.5)

    build_title_page(doc, 6, "Разработка приложения с использованием Hibernate")

    # 1. Текст и вариант задания
    add_heading(doc, "1. Текст и вариант задания", level=1)
    add_p(doc, "Вариант задания: 1 - Книжный магазин либо библиотека.", first_indent=0, space_after=4)
    add_p(doc, "Задание на лабораторную работу:", bold=True, first_indent=0, space_after=4)
    add_bullet(doc, "Используйте в качестве основы лабораторную работу номер 4 или 5 (для заочников – 3).")
    add_bullet(doc, "Добавьте базу данных, модель данных, JPA репозитории и сервисы. Должно быть не менее двух таблиц и одного внешнего ключа.")
    add_bullet(doc, "Реализуйте операции чтения/записи с помощью JPA (spring-boot-data-jpa, также допустимо spring-boot-data-jdbc если сумеете объяснить разницу).")
    add_p(doc, "", first_indent=0, space_after=6)

    # 2. Описание разрабатываемого продукта
    add_heading(doc, "2. Описание разрабатываемого продукта", level=1)
    add_p(doc, "В рамках лабораторной работы №6 веб-приложение «Книжный магазин» переведено с временного in-memory хранилища на полнофункциональную реляционную базу данных под управлением объектно-реляционного отображения Hibernate и спецификации Spring Data JPA. При этом полностью сохранены механизмы безопасности, реализованные в лабораторной работе №4: работа по протоколу HTTPS (порт 8443, SSL/TLS сертификат PKCS12), форма аутентификации со стилизацией ГУАП, разделение ролей USER/ADMIN, защита изменяющих эндпоинтов, CSRF-токены и аудит действий.")

    add_heading(doc, "2.1. Реляционная модель данных и связи сущностей", level=2)
    add_p(doc, "В соответствии с требованиями задания в базу данных добавлены две взаимосвязанные таблицы, объединенные отношением внешнего ключа (Foreign Key):")
    add_bullet(doc, "Таблица «publishers» (сущность Publisher) — представляет издательства книг. Содержит поля: id (первичный ключ, автоинкремент), name (наименование издательства, VARCHAR(100)), city (город расположения, VARCHAR(100)), contact_phone (контактный телефон, VARCHAR(50)).")
    add_bullet(doc, "Таблица «books» (сущность Book) — представляет каталог книг книжного магазина. Содержит поля: id (первичный ключ, автоинкремент), title (название книги, VARCHAR(200)), author (автор, VARCHAR(150)), genre (жанр, VARCHAR(100)), pub_year (год издания, INT), price (цена в рублях, DOUBLE PRECISION), stock (остаток на складе, INT), publisher_id (внешний ключ FK, ссылающийся на id таблицы publishers).")
    add_p(doc, "Связь между сущностями организована отношением @ManyToOne со стороны книги (каждая книга выпущена конкретным издательством) и обратным отношением @OneToMany со стороны издательства (одно издательство может выпустить множество книг).")

    add_heading(doc, "2.2. Применение шаблона проектирования DTO (Data Transfer Object)", level=2)
    add_p(doc, "Как подчеркивается в методических указаниях к лабораторной работе, прямое использование классов JPA-сущностей во внешнем интерфейсе API является антипаттерном из-за риска передачи на клиент избыточных данных, раскрытия внутренней структуры базы данных и потенциального падения приложения при циклических связях. Для решения этой проблемы реализован паттерн DTO:")
    add_bullet(doc, "BookDto — содержит плоскую структуру данных книги, а также идентификатор (publisherId) и наименование издательства (publisherName) без прямого включения вложенных объектов сущностей.")
    add_bullet(doc, "PublisherDto — объект передачи сведений об издательстве.")
    add_p(doc, "Преобразование между сущностями JPA и DTO выполняется в сервисном слое перед отправкой клиенту и при приеме входящих данных.")

    add_heading(doc, "2.3. Управление транзакциями и репозитории", level=2)
    add_p(doc, "Доступ к базе данных реализован через Spring Data JPA репозитории BookRepository и PublisherRepository, расширяющие интерфейс JpaRepository. Бизнес-логика вынесена в сервисные классы BookService и PublisherService, методы которых аннотированы @Transactional. Для операций чтения используется режим @Transactional(readOnly = true), оптимизирующий работу сессий Hibernate. При возникновении исключений Spring обеспечивает автоматический откат транзакций. Изменяющие операции сопровождаются логированием в AuditService.")

    add_heading(doc, "2.4. Стек технологий", level=2)
    add_bullet(doc, "Java 17 (OpenJDK 17 HotSpot)")
    add_bullet(doc, "Spring Boot 3.2.5 (spring-boot-starter-web, spring-boot-starter-data-jpa)")
    add_bullet(doc, "Hibernate ORM 6.4.4 (реализация JPA стандарта)")
    add_bullet(doc, "H2 Database Engine (встроенная реляционная СУБД с web-консолью /h2-console)")
    add_bullet(doc, "HikariCP (высокопроизводительный пул соединений с базой данных)")
    add_bullet(doc, "Spring Security (HTTPS на 8443, Form Login, роли USER/ADMIN, аудит)")
    add_bullet(doc, "Thymeleaf + AngularJS 1.7.5 (клиентское представление с табами каталога и издательств)")
    add_bullet(doc, "SpringDoc OpenAPI 2.5.0 (документация REST API)")

    # 3. Текст основных фрагментов кода
    add_heading(doc, "3. Текст основных фрагментов кода", level=1)

    add_heading(doc, "3.1. Конфигурация зависимостей в pom.xml (фрагмент JPA и H2)", level=2)
    add_code(doc, """<!-- Spring Boot Data JPA & Hibernate -->
<dependency>
    <groupId>org.springframework.boot</groupId>
    <artifactId>spring-boot-starter-data-jpa</artifactId>
</dependency>

<!-- Встроенная СУБД H2 Database -->
<dependency>
    <groupId>com.h2database</groupId>
    <artifactId>h2</artifactId>
    <scope>runtime</scope>
</dependency>""")

    add_heading(doc, "3.2. Настройки источника данных и Hibernate в application.properties", level=2)
    add_code(doc, """server.port=8443
spring.application.name=bookstore

# SSL / HTTPS (порт 8443)
server.ssl.key-store=classpath:keystore.p12
server.ssl.key-store-password=mypassword
server.ssl.keyStoreType=PKCS12
server.ssl.keyAlias=tomcat

# Настройки подключения к СУБД H2 и JPA/Hibernate
spring.datasource.url=jdbc:h2:mem:bookstoredb;DB_CLOSE_DELAY=-1;DB_CLOSE_ON_EXIT=FALSE
spring.datasource.driver-class-name=org.h2.Driver
spring.datasource.username=sa
spring.datasource.password=
spring.jpa.database-platform=org.hibernate.dialect.H2Dialect
spring.jpa.hibernate.ddl-auto=update
spring.jpa.show-sql=false
spring.jpa.properties.hibernate.format_sql=true
spring.h2.console.enabled=true
spring.h2.console.path=/h2-console

# Аудит и логирование
logging.level.ru.spbguap.bookstore=DEBUG
logging.level.org.springframework.security=INFO""")

    add_heading(doc, "3.3. Сущность издательства Publisher.java (первая таблица)", level=2)
    add_code(doc, """@Entity
@Table(name = "publishers")
@Getter
@Setter
@NoArgsConstructor
@AllArgsConstructor
@Builder
public class Publisher {

    @Id
    @GeneratedValue(strategy = GenerationType.IDENTITY)
    private Long id;

    @Column(nullable = false, length = 100)
    private String name;

    @Column(length = 100)
    private String city;

    @Column(length = 50)
    private String contactPhone;

    @OneToMany(mappedBy = "publisher", cascade = CascadeType.ALL, orphanRemoval = false)
    @Builder.Default
    private List<Book> books = new ArrayList<>();
}""")

    add_heading(doc, "3.4. Сущность книги Book.java со связью @ManyToOne (вторая таблица и FK)", level=2)
    add_code(doc, """@Entity
@Table(name = "books")
@Getter
@Setter
@NoArgsConstructor
@AllArgsConstructor
@Builder
public class Book {

    @Id
    @GeneratedValue(strategy = GenerationType.IDENTITY)
    private Long id;

    @Column(nullable = false, length = 200)
    private String title;

    @Column(nullable = false, length = 150)
    private String author;

    @Column(length = 100)
    private String genre;

    @Column(name = "pub_year")
    private Integer year;

    @Column(nullable = false)
    private Double price;

    @Column(nullable = false)
    private Integer stock;

    @ManyToOne(fetch = FetchType.EAGER)
    @JoinColumn(name = "publisher_id")
    private Publisher publisher;
}""")

    add_heading(doc, "3.5. Шаблон Data Transfer Object (BookDto.java)", level=2)
    add_code(doc, """@Data
@Builder
@NoArgsConstructor
@AllArgsConstructor
public class BookDto {
    private Long id;
    private String title;
    private String author;
    private String genre;
    private Integer year;
    private Double price;
    private Integer stock;
    private Long publisherId;
    private String publisherName;
}""")

    add_heading(doc, "3.6. Spring Data JPA репозитории (BookRepository и PublisherRepository)", level=2)
    add_code(doc, """@Repository
public interface BookRepository extends JpaRepository<Book, Long> {
    List<Book> findByAuthorContainingIgnoreCase(String author);
    List<Book> findByGenreIgnoreCase(String genre);
    List<Book> findByPublisherId(Long publisherId);
}

@Repository
public interface PublisherRepository extends JpaRepository<Publisher, Long> {
    Optional<Publisher> findByName(String name);
}""")

    add_heading(doc, "3.7. Сервисный слой с поддержкой транзакций (BookService.java)", level=2)
    add_code(doc, """@Service
@RequiredArgsConstructor
@Transactional
public class BookService {

    private final BookRepository bookRepository;
    private final PublisherRepository publisherRepository;
    private final AuditService auditService;

    public BookDto create(BookDto dto) {
        Book book = toEntity(dto);
        Book saved = bookRepository.save(book);
        auditService.logCreate("Book", saved.getId());
        return toDto(saved);
    }

    @Transactional(readOnly = true)
    public List<BookDto> getAll() {
        return bookRepository.findAll().stream().map(this::toDto).collect(Collectors.toList());
    }

    @Transactional(readOnly = true)
    public Optional<BookDto> getById(Long id) {
        return bookRepository.findById(id).map(this::toDto);
    }

    public Optional<BookDto> update(Long id, BookDto dto) {
        return bookRepository.findById(id).map(existing -> {
            existing.setTitle(dto.getTitle());
            existing.setAuthor(dto.getAuthor());
            existing.setGenre(dto.getGenre());
            existing.setYear(dto.getYear());
            existing.setPrice(dto.getPrice());
            existing.setStock(dto.getStock());
            if (dto.getPublisherId() != null) {
                Publisher publisher = publisherRepository.findById(dto.getPublisherId()).orElse(null);
                existing.setPublisher(publisher);
            }
            Book updated = bookRepository.save(existing);
            auditService.logUpdate("Book", updated.getId());
            return toDto(updated);
        });
    }

    public boolean delete(Long id) {
        if (bookRepository.existsById(id)) {
            bookRepository.deleteById(id);
            auditService.logDelete("Book", id, true);
            return true;
        }
        auditService.logDelete("Book", id, false);
        return false;
    }
}""")

    add_heading(doc, "3.8. Контроллер REST API для управления книгами (BookRestController.java)", level=2)
    add_code(doc, """@RestController
@RequestMapping("/api/books")
@RequiredArgsConstructor
@Tag(name = "Книги", description = "API для управления каталогом книг (JPA / Hibernate)")
public class BookRestController {

    private final BookService bookService;

    @PostMapping
    public ResponseEntity<BookDto> createBook(@RequestBody BookDto bookDto) {
        BookDto created = bookService.create(bookDto);
        return ResponseEntity.status(HttpStatus.CREATED).body(created);
    }

    @GetMapping
    public ResponseEntity<List<BookDto>> getAllBooks() {
        return ResponseEntity.ok(bookService.getAll());
    }

    @GetMapping("/{id}")
    public ResponseEntity<BookDto> getBookById(@PathVariable Long id) {
        return bookService.getById(id)
                .map(ResponseEntity::ok)
                .orElse(ResponseEntity.notFound().build());
    }

    @PutMapping("/{id}")
    public ResponseEntity<BookDto> updateBook(@PathVariable Long id, @RequestBody BookDto bookDto) {
        return bookService.update(id, bookDto)
                .map(ResponseEntity::ok)
                .orElse(ResponseEntity.notFound().build());
    }

    @DeleteMapping("/{id}")
    public ResponseEntity<Void> deleteBook(@PathVariable Long id) {
        if (bookService.delete(id)) {
            return ResponseEntity.noContent().build();
        }
        return ResponseEntity.notFound().build();
    }
}""")

    add_heading(doc, "3.9. Первичное наполнение данными (DataInitializer.java)", level=2)
    add_code(doc, """@Component
@RequiredArgsConstructor
@Slf4j
public class DataInitializer implements CommandLineRunner {

    private final PublisherRepository publisherRepository;
    private final BookRepository bookRepository;

    @Override
    public void run(String... args) {
        if (publisherRepository.count() == 0) {
            Publisher ast = publisherRepository.save(Publisher.builder()
                    .name("АСТ").city("Москва").contactPhone("+7 (495) 789-45-00").build());
            Publisher eksmo = publisherRepository.save(Publisher.builder()
                    .name("Эксмо").city("Москва").contactPhone("+7 (495) 411-68-86").build());
            Publisher azbuka = publisherRepository.save(Publisher.builder()
                    .name("Азбука-Аттикус").city("Санкт-Петербург").contactPhone("+7 (812) 327-04-55").build());

            bookRepository.save(Book.builder()
                    .title("Война и мир").author("Лев Толстой").genre("Роман")
                    .year(1869).price(590.0).stock(5).publisher(ast).build());
            bookRepository.save(Book.builder()
                    .title("Преступление и наказание").author("Фёдор Достоевский").genre("Роман")
                    .year(1866).price(450.0).stock(8).publisher(eksmo).build());
            bookRepository.save(Book.builder()
                    .title("Мастер и Маргарита").author("Михаил Булгаков").genre("Фантастика")
                    .year(1967).price(520.0).stock(12).publisher(azbuka).build());
        }
    }
}""")

    # 4. Вывод
    add_heading(doc, "4. Вывод", level=1)
    add_p(doc, "В ходе выполнения лабораторной работы №6 веб-сервис «Книжный магазин» был успешно переведён на использование объектно-реляционного отображения Hibernate и технологии Spring Data JPA. Спроектирована и реализована реляционная схема базы данных из двух взаимосвязанных таблиц (publishers и books), объединённых отношением внешнего ключа (@ManyToOne и @OneToMany). Реализованы репозитории доступа к данным, сервисный слой с декларативным управлением транзакциями (@Transactional) и поддержкой аудита, а также применён архитектурный паттерн DTO для предотвращения циклических зависимостей и изоляции модели базы данных от внешнего REST API. Все автоматические интеграционные тесты успешно пройдены.")

    doc.save("Отчет_лаб6.docx")
    print("Отчет_лаб6.docx сгенерирован")

# ==============================================================================
# СБОРКА ОТЧЕТА №8
# ==============================================================================
def generate_report_8():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(1.5)

    build_title_page(doc, 8, "Разработка микросервиса")

    # 1. Текст и вариант задания
    add_heading(doc, "1. Текст и вариант задания", level=1)
    add_p(doc, "Вариант задания: 1 - Книжный магазин либо библиотека.", first_indent=0, space_after=4)
    add_p(doc, "Задание на лабораторную работу:", bold=True, first_indent=0, space_after=4)
    add_bullet(doc, "Подготовьте Ваше приложение к разворачиванию в облачном сервисе или компоненте Docker.")
    add_bullet(doc, "Реализуйте файл docker-compose.yml, который будет содержать все необходимые для работы Вашего приложения ресурсы.")
    add_bullet(doc, "Внимание. В виду того, что далеко не на каждом компьютере можно запустить систему виртуализации, данная лабораторная работа сдается в электронном виде, без демонстрации преподавателю.")
    add_p(doc, "", first_indent=0, space_after=6)

    # 2. Описание разрабатываемого продукта
    add_heading(doc, "2. Описание разрабатываемого продукта", level=1)
    add_p(doc, "В рамках лабораторной работы №8 серверное веб-приложение «Книжный магазин» подготовлено к микросервисному контейнерному развертыванию с использованием платформы Docker и системы оркестрации Docker Compose. Создана изолированная мультиконтейнерная архитектура, включающая микросервис приложения и независимый контейнер промышленной СУБД PostgreSQL.")

    add_heading(doc, "2.1. Концепция контейнеризации и архитектура микросервиса", level=2)
    add_p(doc, "В отличие от традиционной аппаратной виртуализации на основе виртуальных машин (где для каждого экземпляра запускается полная гостевая операционная система), контейнеризация Docker использует механизм софтверной изоляции процессов на уровне ядра (пространства имен namespaces и контрольные группы cgroups). Это обеспечивает минимальное потребление системных ресурсов, мгновенный запуск и полную переносимость артефактов между средами разработки и промышленной эксплуатации.")

    add_heading(doc, "2.2. Состав и взаимодействие сервисов в docker-compose.yml", level=2)
    add_p(doc, "Развернутая микросервисная среда включает следующие взаимодействующие компоненты:")
    add_bullet(doc, "Сервис bookstore-app (контейнер приложения) — собирается на базе специализированного манифеста Dockerfile с использованием легковесного образа eclipse-temurin:17-jre-alpine. Приложение запускается по защищенному протоколу HTTPS и публикует порт 8443 наружу на хост-машину. Параметры подключения к базе данных передаются через переменные окружения SPRING_DATASOURCE_*.")
    add_bullet(doc, "Сервис postgres (контейнер реляционной СУБД) — официальный образ postgres:15-alpine. База данных инициализируется скриптом init-db.sql через точку монтирования /docker-entrypoint-initdb.d/ с автоматическим созданием таблиц publishers, books и начальных записей. Для проверки готовности СУБД настроен healthcheck через утилиту pg_isready.")
    add_bullet(doc, "Изолированная сеть bookstore-net (тип bridge) — обеспечивает приватное сетевое взаимодействие между контейнером приложения и контейнером базы данных по внутренним DNS-именам сервисов.")
    add_bullet(doc, "Постоянный том pgdata (persistent volume) — монтируется в каталог /var/lib/postgresql/data СУБД, гарантируя сохранность данных даже при полной остановке и пересоздании контейнеров.")

    add_heading(doc, "2.3. Стек технологий", level=2)
    add_bullet(doc, "Docker Engine 29.3.1 (платформа контейнеризации)")
    add_bullet(doc, "Docker Compose v5.1.1 (декларативная оркестрация мультиконтейнерных сред)")
    add_bullet(doc, "Базовый образ Eclipse Temurin OpenJDK 17 (JRE Alpine Linux)")
    add_bullet(doc, "СУБД PostgreSQL 15 Alpine")
    add_bullet(doc, "Spring Boot 3.2.5 (микросервис на Java 17, Hibernate ORM, Spring Security SSL/HTTPS)")
    add_bullet(doc, "Apache Maven 3.9.16 (сборка JAR-артефакта)")

    # 3. Текст основных фрагментов кода
    add_heading(doc, "3. Текст основных фрагментов кода", level=1)

    add_heading(doc, "3.1. Манифест сборки Dockerfile", level=2)
    add_code(doc, """# Базовый образ: легковесная среда исполнения Eclipse Temurin OpenJDK 17 (JRE Alpine)
FROM eclipse-temurin:17-jre-alpine

LABEL maintainer="K. A. Engelgardt <spbguap@example.com>"
LABEL version="1.0.0"
LABEL description="Bookstore Microservice — TRSIS Lab 8 (SPb GUAP)"

WORKDIR /app

# Копирование предварительно собранного Spring Boot JAR-артефакта
COPY target/bookstore-1.0.0.jar app.jar

# Публикация защищенного порта HTTPS (8443)
EXPOSE 8443

# Настройка виртуальной машины Java
ENV JAVA_OPTS="-Xms256m -Xmx512m -Dfile.encoding=UTF-8"

# Точка входа для запуска микросервиса
ENTRYPOINT ["sh", "-c", "java $JAVA_OPTS -jar app.jar"]""")

    add_heading(doc, "3.2. Файл исключений .dockerignore", level=2)
    add_code(doc, """target/classes
target/test-classes
target/generated-sources
.idea
.git
*.bat
*.log
*.pdf
*.docx""")

    add_heading(doc, "3.3. Спецификация оркестрации docker-compose.yml", level=2)
    add_code(doc, """version: "3.9"

services:
  # ── Сервис СУБД: Реляционная база данных PostgreSQL ──
  postgres:
    image: postgres:15-alpine
    container_name: bookstore-postgres
    environment:
      POSTGRES_DB: bookstore
      POSTGRES_USER: bookstore_user
      POSTGRES_PASSWORD: bookstore_pass
    volumes:
      - ./init-db.sql:/docker-entrypoint-initdb.d/01-init.sql:ro
      - pgdata:/var/lib/postgresql/data
    ports:
      - "5432:5432"
    networks:
      - bookstore-net
    healthcheck:
      test: ["CMD-SHELL", "pg_isready -U bookstore_user -d bookstore"]
      interval: 5s
      timeout: 5s
      retries: 5
    restart: unless-stopped

  # ── Сервис микросервиса: Приложение «Книжный магазин» ──
  bookstore-app:
    build:
      context: .
      dockerfile: Dockerfile
    container_name: bookstore-service
    ports:
      - "8443:8443"
    depends_on:
      postgres:
        condition: service_healthy
    environment:
      SPRING_DATASOURCE_URL: jdbc:postgresql://postgres:5432/bookstore
      SPRING_DATASOURCE_USERNAME: bookstore_user
      SPRING_DATASOURCE_PASSWORD: bookstore_pass
      SPRING_DATASOURCE_DRIVER_CLASS_NAME: org.postgresql.Driver
      SPRING_JPA_DATABASE_PLATFORM: org.hibernate.dialect.PostgreSQLDialect
      SPRING_JPA_HIBERNATE_DDL_AUTO: update
      SPRING_H2_CONSOLE_ENABLED: "false"
    networks:
      - bookstore-net
    restart: unless-stopped

volumes:
  pgdata:
    driver: local

networks:
  bookstore-net:
    driver: bridge""")

    add_heading(doc, "3.4. Скрипт первичной инициализации init-db.sql", level=2)
    add_code(doc, """-- Таблица 1: Издательства (Publishers)
CREATE TABLE IF NOT EXISTS publishers (
    id BIGSERIAL PRIMARY KEY,
    name VARCHAR(100) NOT NULL,
    city VARCHAR(100),
    contact_phone VARCHAR(50)
);

-- Таблица 2: Книги (Books) со связью по внешнему ключу publisher_id
CREATE TABLE IF NOT EXISTS books (
    id BIGSERIAL PRIMARY KEY,
    title VARCHAR(200) NOT NULL,
    author VARCHAR(150) NOT NULL,
    genre VARCHAR(100),
    pub_year INT,
    price DOUBLE PRECISION NOT NULL,
    stock INT NOT NULL,
    publisher_id BIGINT REFERENCES publishers(id) ON DELETE SET NULL
);

-- Первичное наполнение тестовыми данными
INSERT INTO publishers (name, city, contact_phone) VALUES
    ('АСТ', 'Москва', '+7 (495) 789-45-00'),
    ('Эксмо', 'Москва', '+7 (495) 411-68-86'),
    ('Азбука-Аттикус', 'Санкт-Петербург', '+7 (812) 327-04-55')
ON CONFLICT DO NOTHING;

INSERT INTO books (title, author, genre, pub_year, price, stock, publisher_id) VALUES
    ('Война и мир', 'Лев Толстой', 'Роман', 1869, 590.0, 5, 1),
    ('Преступление и наказание', 'Фёдор Достоевский', 'Роман', 1866, 450.0, 8, 2),
    ('Мастер и Маргарита', 'Михаил Булгаков', 'Фантастика', 1967, 520.0, 12, 3)
ON CONFLICT DO NOTHING;""")

    add_heading(doc, "3.5. Команды управления микросервисным окружением", level=2)
    add_code(doc, """# 1. Предварительная сборка JAR-артефакта в Maven:
mvn clean package -DskipTests

# 2. Сборка Docker-образов и запуск мультиконтейнерного комплекса:
docker compose up --build -d

# 3. Проверка статуса запущенных контейнеров:
docker compose ps

# 4. Просмотр логов микросервиса:
docker compose logs -f bookstore-app

# 5. Остановка комплекса с сохранением томов:
docker compose down""")

    # 4. Вывод
    add_heading(doc, "4. Вывод", level=1)
    add_p(doc, "В ходе выполнения лабораторной работы №8 серверное приложение «Книжный магазин» было успешно преобразовано в контейнеризированный микросервис. Разработан легковесный образ Dockerfile на базе Alpine Linux и Eclipse Temurin 17 JRE. Создана мультиконтейнерная конфигурация docker-compose.yml, объединяющая микросервис и реляционную СУБД PostgreSQL 15 в единую изолированную виртуальную сеть с персистентным дисковым томом для сохранения данных. Проверена валидность конфигурации Docker Compose. Работа подготовлена к электронной сдаче согласно требованиям методических указаний.")

    doc.save("Отчет_лаб8.docx")
    print("Отчет_лаб8.docx сгенерирован")

if __name__ == "__main__":
    generate_report_6()
    generate_report_8()

    print("Конвертация в PDF...")
    convert("Отчет_лаб6.docx", "Отчет_лаб6.pdf")
    convert("Отчет_лаб8.docx", "Отчет_лаб8.pdf")
    print("Все отчеты успешно сгенерированы и конвертированы в PDF!")
